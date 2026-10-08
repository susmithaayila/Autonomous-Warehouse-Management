import os
import shutil
import openpyxl

BASE_DIR = r"C:\Users\ayila\.gemini\antigravity-ide\scratch\AWMS"
DIAGRAMS_DIR = os.path.join(BASE_DIR, "diagrams")
EXCEL_DIR = os.path.join(BASE_DIR, "excel")

# 1. Convert Excel files to Markdown for easy viewing in IDE
def convert_excel_to_md(xlsx_path, md_path):
    if not os.path.exists(xlsx_path) or os.path.basename(xlsx_path).startswith("~$"):
        return
    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        with open(md_path, "w", encoding="utf-8") as f:
            f.write(f"# {os.path.basename(xlsx_path).replace('.xlsx', '')} Data View\n\n")
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                f.write(f"## Sheet: {sheet_name}\n\n")
                rows = list(ws.iter_rows(values_only=True))
                if not rows:
                    continue
                headers = [str(cell or '') for cell in rows[0]]
                f.write("| " + " | ".join(headers) + " |\n")
                f.write("| " + " | ".join(["---"] * len(headers)) + " |\n")
                for row in rows[1:]:
                    row_vals = [str(cell if cell is not None else '') for cell in row]
                    f.write("| " + " | ".join(row_vals) + " |\n")
                f.write("\n\n")
        print(f"Created MD view: {os.path.basename(md_path)}")
    except Exception as e:
        print(f"Skipping {xlsx_path}: {e}")

# Convert all Excels in excel/ and root
for xlsx_name in os.listdir(EXCEL_DIR):
    if xlsx_name.endswith(".xlsx") and not xlsx_name.startswith("~$"):
        xlsx_p = os.path.join(EXCEL_DIR, xlsx_name)
        md_p = os.path.join(EXCEL_DIR, xlsx_name.replace(".xlsx", ".md"))
        convert_excel_to_md(xlsx_p, md_p)

# Root Traceability Matrix
convert_excel_to_md(os.path.join(BASE_DIR, "Traceability_Matrix.xlsx"), os.path.join(BASE_DIR, "Traceability_Matrix.md"))

# 2. Convert Draw.io files to standalone viewer HTML pages
def convert_drawio_to_html(drawio_path, html_path, title):
    if not os.path.exists(drawio_path):
        return
    with open(drawio_path, "r", encoding="utf-8") as f:
        xml_content = f.read()
    
    html_content = f"""<!DOCTYPE html>
<html>
<head>
    <title>{title}</title>
    <script type="text/javascript" src="https://viewer.diagrams.net/js/viewer-static.min.js"></script>
    <style>
        body {{ margin: 0; padding: 20px; background: #0f172a; color: #fff; font-family: sans-serif; text-align: center; }}
        h1 {{ color: #60a5fa; margin-bottom: 20px; }}
        .mxgraph {{ background: #ffffff; border-radius: 12px; padding: 20px; display: inline-block; box-shadow: 0 10px 25px rgba(0,0,0,0.5); }}
    </style>
</head>
<body>
    <h1>🎨 {title}</h1>
    <div class="mxgraph" style="max-width:100%;border:1px solid #ccc;" data-mxgraph="{{&quot;highlight&quot;:&quot;#0000ff&quot;,&quot;nav&quot;:true,&quot;resize&quot;:true,&quot;toolbar&quot;:&quot;zoom layers tags lightbox&quot;,&quot;xml&quot;:&quot;{xml_content.replace('"', '&quot;').replace('\n', '')}&quot;}}"></div>
</body>
</html>"""
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html_content)
    print(f"Created HTML diagram view: {os.path.basename(html_path)}")

diagram_titles = {
    "P03_Use_Case_Diagram.drawio": "AWMS Use Case Diagram",
    "P03_Analysis_Model.drawio": "AWMS Scenario Analysis Model",
    "P04_ER_Diagram.drawio": "AWMS Entity Relationship Diagram (ERD)",
    "P04_DFD_Level_0.drawio": "AWMS Level 0 DFD",
    "P04_DFD_Level_1.drawio": "AWMS Level 1 DFD with Trust Boundaries",
    "P05_Architecture.drawio": "AWMS Secure Layered Architecture",
    "P07_STRIDE_DFD.drawio": "AWMS STRIDE Threat Model DFD",
    "P08_Attack_Tree.drawio": "AWMS Security Attack Tree",
    "P08_Security_Refined_Architecture.drawio": "AWMS Security Refined Architecture"
}

for d_name, d_title in diagram_titles.items():
    d_path = os.path.join(DIAGRAMS_DIR, d_name)
    h_path = os.path.join(DIAGRAMS_DIR, d_name.replace(".drawio", ".html"))
    convert_drawio_to_html(d_path, h_path, d_title)

# Also copy all .html diagram viewers and .md table viewers to evidence folders
evidence_dir = os.path.join(BASE_DIR, "evidence")
for d_name, d_title in diagram_titles.items():
    h_src = os.path.join(DIAGRAMS_DIR, d_name.replace(".drawio", ".html"))
    if "P03" in d_name:
        shutil.copy2(h_src, os.path.join(evidence_dir, "P03_UML", os.path.basename(h_src)))
    elif "P04" in d_name:
        shutil.copy2(h_src, os.path.join(evidence_dir, "P04_DataFlow", os.path.basename(h_src)))
    elif "P05" in d_name:
        shutil.copy2(h_src, os.path.join(evidence_dir, "P05_Architecture", os.path.basename(h_src)))
    elif "P07" in d_name:
        shutil.copy2(h_src, os.path.join(evidence_dir, "P07_ThreatModel", os.path.basename(h_src)))
    elif "P08" in d_name:
        shutil.copy2(h_src, os.path.join(evidence_dir, "P08_AttackTree", os.path.basename(h_src)))

for xlsx_name in os.listdir(EXCEL_DIR):
    if xlsx_name.endswith(".md"):
        m_src = os.path.join(EXCEL_DIR, xlsx_name)
        # copy to corresponding evidence folder
        p_name = xlsx_name.split("_")[0]
        for f_name in os.listdir(evidence_dir):
            if f_name.startswith(p_name):
                shutil.copy2(m_src, os.path.join(evidence_dir, f_name, xlsx_name))

print("Conversion and copy complete!")
